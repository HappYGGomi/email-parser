"""
Excel Exporter Module - 질의응답 데이터를 엑셀 파일로 저장
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict
from datetime import datetime


class ExcelExporter:
    """엑셀 파일 생성 클래스"""

    def __init__(self):
        self.wb = None
        self.ws = None

    def export_qa_pairs(self, qa_pairs: List[Dict], output_file: str):
        """질의응답 쌍을 엑셀로 저장"""
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "이메일 질의응답"

        # 헤더 작성
        self._write_header()

        # 데이터 작성
        self._write_data(qa_pairs)

        # 스타일 적용
        self._apply_styles()

        # 파일 저장
        self.wb.save(output_file)

        return output_file

    def _write_header(self):
        """헤더 작성"""
        headers = [
            '번호',
            '스레드 제목',
            '질문 발신자',
            '질문 일시',
            '질문 내용',
            '답변 발신자',
            '답변 일시',
            '답변 내용'
        ]

        for col, header in enumerate(headers, start=1):
            cell = self.ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, size=11, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def _write_data(self, qa_pairs: List[Dict]):
        """데이터 작성"""
        for idx, qa in enumerate(qa_pairs, start=1):
            row = idx + 1

            # 번호
            self.ws.cell(row=row, column=1).value = idx

            # 스레드 제목
            self.ws.cell(row=row, column=2).value = qa.get('thread_subject', '')

            # 질문 발신자
            self.ws.cell(row=row, column=3).value = qa.get('question_sender', '')

            # 질문 일시
            self.ws.cell(row=row, column=4).value = qa.get('question_date', '')

            # 질문 내용
            question_cell = self.ws.cell(row=row, column=5)
            question_cell.value = qa.get('question', '')
            question_cell.alignment = Alignment(wrap_text=True, vertical='top')

            # 답변 발신자
            self.ws.cell(row=row, column=6).value = qa.get('answer_sender', '')

            # 답변 일시
            self.ws.cell(row=row, column=7).value = qa.get('answer_date', '')

            # 답변 내용
            answer_cell = self.ws.cell(row=row, column=8)
            answer_cell.value = qa.get('answer', '')
            answer_cell.alignment = Alignment(wrap_text=True, vertical='top')

    def _apply_styles(self):
        """스타일 적용"""
        # 테두리 스타일
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 모든 셀에 테두리 적용
        for row in self.ws.iter_rows(min_row=1, max_row=self.ws.max_row, min_col=1, max_col=8):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:  # 헤더가 아닌 경우
                    cell.alignment = Alignment(vertical='top', wrap_text=True)

        # 컬럼 너비 조정
        column_widths = {
            'A': 8,   # 번호
            'B': 30,  # 스레드 제목
            'C': 25,  # 질문 발신자
            'D': 20,  # 질문 일시
            'E': 50,  # 질문 내용
            'F': 25,  # 답변 발신자
            'G': 20,  # 답변 일시
            'H': 50   # 답변 내용
        }

        for col, width in column_widths.items():
            self.ws.column_dimensions[col].width = width

        # 행 높이 자동 조정
        for row in range(2, self.ws.max_row + 1):
            self.ws.row_dimensions[row].height = None  # 자동 높이

        # 첫 행 고정
        self.ws.freeze_panes = 'A2'

    def export_raw_emails(self, messages: List[Dict], output_file: str):
        """원본 이메일 데이터를 엑셀로 저장"""
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "원본 이메일"

        # 헤더 작성
        headers = ['번호', '제목', '발신자', '수신자', '일시', '본문']
        for col, header in enumerate(headers, start=1):
            cell = self.ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, size=11, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # 데이터 작성
        for idx, msg in enumerate(messages, start=1):
            row = idx + 1
            self.ws.cell(row=row, column=1).value = idx
            self.ws.cell(row=row, column=2).value = msg.get('subject', '')
            self.ws.cell(row=row, column=3).value = msg.get('sender', '')
            self.ws.cell(row=row, column=4).value = msg.get('recipients', '')
            self.ws.cell(row=row, column=5).value = msg.get('date', '')

            body_cell = self.ws.cell(row=row, column=6)
            body_cell.value = msg.get('body', '')
            body_cell.alignment = Alignment(wrap_text=True, vertical='top')

        # 스타일 적용
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for row in self.ws.iter_rows(min_row=1, max_row=self.ws.max_row, min_col=1, max_col=6):
            for cell in row:
                cell.border = thin_border

        # 컬럼 너비 조정
        self.ws.column_dimensions['A'].width = 8
        self.ws.column_dimensions['B'].width = 35
        self.ws.column_dimensions['C'].width = 30
        self.ws.column_dimensions['D'].width = 30
        self.ws.column_dimensions['E'].width = 20
        self.ws.column_dimensions['F'].width = 60

        # 첫 행 고정
        self.ws.freeze_panes = 'A2'

        # 파일 저장
        self.wb.save(output_file)

        return output_file


if __name__ == '__main__':
    # 테스트 코드
    print("Excel Exporter Module loaded successfully!")
